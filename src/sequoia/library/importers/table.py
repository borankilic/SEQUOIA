import csv
import math
import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Literal

import numpy as np
from pydantic import ValidationError
from scipy import constants

from sequoia.data import Curve, Provenance
from sequoia.errors import TableImportError, format_validation_error
from sequoia.units import from_si, to_si, to_wavelength, unit_dimension

YKind = Literal["fraction", "percent", "od", "db", "responsivity", "raw"]

_TEXT_SUFFIXES = {".csv", ".tsv", ".txt", ".dat", ".asc"}
_EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
_UNIT_PATTERNS = (
    re.compile(r"^(?P<name>.*?)\s*\((?P<unit>[^()]*)\)\s*$"),
    re.compile(r"^(?P<name>.*?)\s*\[(?P<unit>[^\[\]]*)\]\s*$"),
)


@dataclass(frozen=True)
class Column:
    header: str
    name: str
    unit: str | None
    values: np.ndarray


@dataclass(frozen=True)
class Table:
    path: Path
    sheet: str | None
    header_row: int | None
    columns: tuple[Column, ...]

    def column(self, ref: int | str) -> Column:
        if isinstance(ref, int):
            if not -len(self.columns) <= ref < len(self.columns):
                raise TableImportError(
                    f"{self.path.name}: column index {ref} out of range "
                    f"(table has {len(self.columns)} numeric columns)"
                )
            return self.columns[ref]
        wanted = ref.strip().lower()
        matches = [
            c for c in self.columns if wanted in (c.header.strip().lower(), c.name.strip().lower())
        ]
        if len(matches) != 1:
            headers = [c.header for c in self.columns]
            problem = "no column" if not matches else "more than one column"
            raise TableImportError(
                f"{self.path.name}: {problem} matches {ref!r}; columns: {headers}"
            )
        return matches[0]


@dataclass(frozen=True)
class ImportResult:
    curve: Curve
    warnings: tuple[str, ...]


def _normalize_unit(unit: str) -> str:
    u = unit.strip().replace("µ", "u").replace("μ", "u")
    if re.fullmatch(r"cm\^?-1|1/cm|cm⁻¹", u, re.IGNORECASE):
        return "1/cm"
    if u.lower() in ("%t", "%", "percent"):
        return "%"
    if u == "Å":
        return "angstrom"
    return u


def _split_name_unit(header: str) -> tuple[str, str | None]:
    text = header.strip()
    for pattern in _UNIT_PATTERNS:
        match = pattern.match(text)
        if match and match["unit"].strip():
            return match["name"].strip(), _normalize_unit(match["unit"])
    return text, None


def _to_float(cell: Any, decimal_comma: bool) -> float:
    if isinstance(cell, bool) or cell is None:
        return math.nan
    if isinstance(cell, int | float):
        return float(cell)
    text = str(cell).strip()
    if decimal_comma:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return math.nan


def _cell_text(cell: Any) -> str:
    return "" if cell is None else str(cell).strip()


def _split_line(line: str, delimiter: str | None) -> list[str]:
    if delimiter is None:
        return line.split()
    return [c.strip() for c in next(csv.reader([line], delimiter=delimiter))]


def _detect_delimiter(lines: list[str]) -> str | None:
    sample = lines[:200]
    for delimiter in ("\t", ";", ",", None):
        counts = Counter(len(_split_line(line, delimiter)) for line in sample)
        width, freq = counts.most_common(1)[0]
        if width >= 2 and freq / len(sample) >= 0.6:
            return delimiter
    return None


def _read_text(path: Path) -> tuple[list[list[Any]], bool, bool]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        raise TableImportError(f"{path.name}: file is empty")
    delimiter = _detect_delimiter(lines)
    decimal_comma = delimiter != "," and any(re.search(r"\d,\d", line) for line in lines)
    grid = [_split_line(line, delimiter) for line in lines]
    return grid, decimal_comma, delimiter is None


def list_sheets(path: Path) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _read_excel(path: Path, sheet: str | None) -> list[list[Any]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise TableImportError(f"{path.name}: cannot open workbook: {exc}") from exc
    try:
        if sheet is None:
            ws = wb.worksheets[0]
        elif sheet in wb.sheetnames:
            ws = wb[sheet]
        else:
            raise TableImportError(f"{path.name}: no sheet {sheet!r}; sheets: {wb.sheetnames}")
        return [list(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _grid_to_table(
    grid: list[list[Any]], path: Path, sheet: str | None, decimal_comma: bool, ragged_header: bool
) -> Table:
    nums = [[_to_float(c, decimal_comma) for c in row] for row in grid]
    counts = [sum(not math.isnan(v) for v in row) for row in nums]
    last = len(grid) - 1
    start = next(
        (i for i in range(len(grid)) if counts[i] >= 2 and (i == last or counts[i + 1] >= 1)), None
    )
    if start is None:
        raise TableImportError(
            f"{path.name}: no numeric data found (need at least two numeric columns)"
        )
    end = start
    while end <= last and counts[end] >= 1:
        end += 1
    cols = [j for j, v in enumerate(nums[start]) if not math.isnan(v)]
    data = np.array(
        [[nums[i][j] if j < len(nums[i]) else math.nan for j in cols] for i in range(start, end)],
        dtype=float,
    )
    header_row = None
    headers = [f"column {k + 1}" for k in range(len(cols))]
    for h in range(start - 1, -1, -1):
        row = grid[h]
        texts = [_cell_text(row[j]) if j < len(row) else "" for j in cols]
        if any(t and math.isnan(_to_float(t, decimal_comma)) for t in texts):
            if not (ragged_header and len(row) != len(grid[start])):
                header_row = h
                headers = [t or f"column {k + 1}" for k, t in enumerate(texts)]
            break
    columns = []
    for k, header in enumerate(headers):
        name, unit = _split_name_unit(header)
        columns.append(Column(header, name or header, unit, data[:, k]))
    return Table(path, sheet, header_row, tuple(columns))


def read_table(path: Path | str, sheet: str | None = None) -> Table:
    path = Path(path)
    if not path.is_file():
        raise TableImportError(f"{path}: file not found")
    suffix = path.suffix.lower()
    if suffix in _EXCEL_SUFFIXES:
        return _grid_to_table(_read_excel(path, sheet), path, sheet, False, False)
    if suffix in _TEXT_SUFFIXES:
        grid, decimal_comma, ragged = _read_text(path)
        return _grid_to_table(grid, path, None, decimal_comma, ragged)
    raise TableImportError(
        f"{path.name}: unsupported file type {suffix!r} "
        f"(use {', '.join(sorted(_TEXT_SUFFIXES | _EXCEL_SUFFIXES))})"
    )


def guess_y_kind(name: str, unit: str | None) -> YKind:
    u = (unit or "").strip().lower()
    n = name.strip().lower()
    if u in ("%", "percent") or n.startswith("%"):
        return "percent"
    if u == "od" or n in ("od", "optical density", "optical depth"):
        return "od"
    if u == "db":
        return "db"
    if u == "a/w":
        return "responsivity"
    if u in ("", "1", "fraction"):
        return "fraction"
    return "raw"


def _to_fraction(values: np.ndarray, kind: str) -> np.ndarray:
    if kind == "percent":
        return values / 100.0
    if kind == "od":
        return 10.0 ** (-values)
    if kind == "db":
        return 10.0 ** (-values / 10.0)
    return values


def _default_y_label(col: Column, kind: str, target: str) -> str:
    if kind != "raw" and target == "responsivity":
        return "responsivity"
    label = col.name.strip().lower().replace(" ", "_")
    return label if label and label not in ("od", "%", "db") else "value"


def import_curve(
    table: Table,
    x: int | str,
    y: int | str,
    *,
    x_unit: str | None = None,
    y_unit: str | None = None,
    y_kind: YKind | None = None,
    target: Literal["fraction", "responsivity"] = "fraction",
    y_label: str | None = None,
    provenance: Provenance | None = None,
) -> ImportResult:
    x_col, y_col = table.column(x), table.column(y)
    notes: list[str] = []

    xu = _normalize_unit(x_unit) if x_unit else x_col.unit
    if xu is None:
        raise TableImportError(
            f"unit of x column {x_col.header!r} is unknown; pass x_unit (e.g. 'nm')"
        )
    try:
        dimension = unit_dimension(xu)
    except ValueError as exc:
        raise TableImportError(f"x column {x_col.header!r}: {exc}") from exc
    kind: YKind = y_kind or guess_y_kind(y_col.name, y_col.unit)

    xv, yv = x_col.values, y_col.values
    keep = np.isfinite(xv) & np.isfinite(yv)
    if not keep.all():
        notes.append(f"dropped {int((~keep).sum())} row(s) with missing or non-numeric values")
        xv, yv = xv[keep], yv[keep]

    wavelength_m = None
    if dimension in ("frequency", "wavenumber", "energy"):
        if np.any(xv <= 0):
            raise TableImportError("x values must be positive to convert them to wavelength")
        wavelength_m = to_wavelength(to_si(xv, xu), dimension)
        x_out, x_out_unit, x_label = from_si(wavelength_m, "nm"), "nm", "wavelength"
    elif dimension == "length":
        wavelength_m = to_si(xv, xu)
        x_out, x_out_unit, x_label = xv, xu, "wavelength"
    else:
        x_out, x_out_unit, x_label = xv, xu, x_col.name.strip().lower() or "x"

    order = np.argsort(x_out, kind="stable")
    x_out, yv = x_out[order], yv[order]
    if wavelength_m is not None:
        wavelength_m = wavelength_m[order]
    duplicates = np.flatnonzero(np.diff(x_out) == 0)
    if duplicates.size:
        raise TableImportError(
            f"duplicate x value {x_out[duplicates[0]]:g} {x_out_unit} "
            f"({duplicates.size} duplicate(s)); remove or average them first"
        )

    hc_over_q = constants.h * constants.c / constants.e
    if kind == "raw":
        y_out = yv
        y_out_unit = _normalize_unit(y_unit) if y_unit else (y_col.unit or "1")
    else:
        if (kind == "responsivity") != (target == "responsivity") and wavelength_m is None:
            raise TableImportError(
                "converting between responsivity and efficiency needs a wavelength x axis"
            )
        if target == "fraction":
            if kind == "responsivity":
                y_out = yv * hc_over_q / wavelength_m
            else:
                y_out = _to_fraction(yv, kind)
            y_out_unit = "1"
        else:
            if kind == "responsivity":
                y_out = yv
            else:
                y_out = _to_fraction(yv, kind) * wavelength_m / hc_over_q
            y_out_unit = "A/W"
        if target == "fraction" and y_out.size and (y_out.min() < -1e-9 or y_out.max() > 1 + 1e-9):
            hint = " (looks like percent; set y_kind='percent')" if kind == "fraction" else ""
            notes.append(
                f"values span {y_out.min():.4g} to {y_out.max():.4g}, outside [0, 1]{hint}"
            )

    try:
        curve = Curve(
            x=x_out.tolist(),
            y=np.asarray(y_out, dtype=float).tolist(),
            x_label=x_label,
            x_unit=x_out_unit,
            y_label=y_label or _default_y_label(y_col, kind, target),
            y_unit=y_out_unit,
            provenance=provenance or Provenance(source="raw-data", reference=table.path.name),
        )
    except ValidationError as exc:
        raise TableImportError(
            f"{table.path.name}: invalid curve:\n{format_validation_error(exc)}"
        ) from exc
    return ImportResult(curve, tuple(notes))
