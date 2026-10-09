import numpy as np
import pytest
from openpyxl import Workbook
from scipy import constants

from sequoia.data import Provenance
from sequoia.errors import TableImportError
from sequoia.library.importers.table import (
    guess_y_kind,
    import_curve,
    list_sheets,
    read_table,
)

WAVELENGTHS = np.array([400.0, 450.0, 500.0, 550.0, 600.0, 650.0, 700.0])
TRANSMISSION = np.array([0.012, 0.2, 0.5, 0.875, 0.95, 0.5, 0.0234])


def write_text(path, lines, encoding="utf-8"):
    path.write_text("\n".join(lines) + "\n", encoding=encoding)
    return path


def percent_csv(tmp_path):
    rows = [
        f"{w!r},{t * 100!r}"
        for w, t in zip(WAVELENGTHS.tolist(), TRANSMISSION.tolist(), strict=True)
    ]
    return write_text(tmp_path / "percent.csv", ["Wavelength (nm),Transmission (%)", *rows])


def od_csv(tmp_path):
    rows = [
        f"{w!r};{float(-np.log10(t))!r}".replace(".", ",")
        for w, t in zip(WAVELENGTHS.tolist(), TRANSMISSION.tolist(), strict=True)
    ]
    return write_text(tmp_path / "od.csv", ["Wavelength [nm];OD", *rows])


def fraction_xlsx(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Filter FB405-10"])
    ws.append([])
    ws.append(["Wavelength (nm)", "Transmission"])
    for w, t in zip(WAVELENGTHS, TRANSMISSION, strict=True):
        ws.append([float(w), float(t)])
    ws.append([])
    ws.append(["Data are typical"])
    path = tmp_path / "fraction.xlsx"
    wb.save(path)
    return path


def test_percent_od_and_xlsx_import_to_identical_curves(tmp_path):
    curves = []
    for path in (percent_csv(tmp_path), od_csv(tmp_path), fraction_xlsx(tmp_path)):
        result = import_curve(read_table(path), 0, 1)
        assert result.warnings == ()
        curves.append(result.curve)
    for curve in curves:
        assert (curve.x_unit, curve.y_unit, curve.x_label) == ("nm", "1", "wavelength")
        np.testing.assert_allclose(curve.x_si, WAVELENGTHS * 1e-9, rtol=1e-12)
        np.testing.assert_allclose(curve.y_si, TRANSMISSION, rtol=1e-12)


def test_header_detection_skips_title_rows_and_footers(tmp_path):
    table = read_table(fraction_xlsx(tmp_path))
    assert table.header_row is not None
    assert [c.name for c in table.columns] == ["Wavelength", "Transmission"]
    assert [c.unit for c in table.columns] == ["nm", None]
    assert len(table.columns[0].values) == len(WAVELENGTHS)


def test_columns_can_be_selected_by_name(tmp_path):
    table = read_table(percent_csv(tmp_path))
    result = import_curve(table, "wavelength", "Transmission (%)")
    assert result.curve.y_label == "transmission"
    with pytest.raises(TableImportError, match="no column matches 'power'"):
        table.column("power")
    with pytest.raises(TableImportError, match="out of range"):
        table.column(5)


def test_whitespace_delimited_ascii_with_multiword_header(tmp_path):
    simple = write_text(
        tmp_path / "simple.txt", ["Wavelength Transmission", "400 0.1", "500 0.9", "600 0.5"]
    )
    assert [c.name for c in read_table(simple).columns] == ["Wavelength", "Transmission"]

    path = write_text(
        tmp_path / "semrock.txt",
        ["Wavelength (nm)   Transmission (%)", "400 10", "500 90", "600 50"],
    )
    table = read_table(path)
    assert [c.header for c in table.columns] == ["column 1", "column 2"]
    result = import_curve(table, 0, 1, x_unit="nm", y_kind="percent")
    assert result.curve.y == [0.1, 0.9, 0.5]


def test_tab_delimited_without_header(tmp_path):
    path = write_text(tmp_path / "plain.tsv", ["400\t0.1", "500\t0.9"])
    table = read_table(path)
    assert [c.header for c in table.columns] == ["column 1", "column 2"]
    with pytest.raises(TableImportError, match="unit of x column"):
        import_curve(table, 0, 1)
    assert import_curve(table, 0, 1, x_unit="nm").curve.x == [400.0, 500.0]


def test_cp1252_micro_sign_in_header(tmp_path):
    path = write_text(
        tmp_path / "tl.csv",
        ["Wavelength (µm),Responsivity (A/W)", "0.4,0.1", "0.8,0.5", "1.0,0.6"],
        encoding="cp1252",
    )
    table = read_table(path)
    assert table.columns[0].unit == "um"
    curve = import_curve(table, 0, 1).curve
    assert curve.x_unit == "um"
    np.testing.assert_allclose(curve.x_si, [0.4e-6, 0.8e-6, 1.0e-6])
    assert curve.y_unit == "1"


def test_frequency_axis_is_converted_to_sorted_wavelength(tmp_path):
    path = write_text(
        tmp_path / "thz.csv",
        ["Frequency (THz),Transmission", "600,0.3", "500,0.9", "400,0.1"],
    )
    curve = import_curve(read_table(path), 0, 1).curve
    assert curve.x_unit == "nm"
    assert curve.y == [0.3, 0.9, 0.1]
    np.testing.assert_allclose(curve.x_si, constants.c / np.array([600e12, 500e12, 400e12]))
    assert np.all(np.diff(curve.x_si) > 0)


def test_wavenumber_axis(tmp_path):
    path = write_text(tmp_path / "k.csv", ["Wavenumber (cm-1),T", "10000,0.1", "20000,0.2"])
    curve = import_curve(read_table(path), 0, 1).curve
    np.testing.assert_allclose(curve.x_si, [500e-9, 1000e-9])
    assert curve.y == [0.2, 0.1]


def test_responsivity_and_quantum_efficiency_are_interconvertible(tmp_path):
    path = write_text(
        tmp_path / "pd.csv",
        ["Wavelength (nm),Responsivity (A/W)", "500,0.25", "800,0.5", "1000,0.6"],
    )
    table = read_table(path)
    qe = import_curve(table, 0, 1, target="fraction").curve
    expected = (
        np.array([0.25, 0.5, 0.6])
        * constants.h
        * constants.c
        / constants.e
        / np.array([500e-9, 800e-9, 1000e-9])
    )
    np.testing.assert_allclose(qe.y_si, expected, rtol=1e-12)
    assert qe.y_unit == "1"
    same = import_curve(table, 0, 1, target="responsivity").curve
    assert same.y == [0.25, 0.5, 0.6]

    percent_path = write_text(
        tmp_path / "qe.csv",
        ["Wavelength (nm),QE (%)", "500,62.1", "800,38.2"],
    )
    back = import_curve(read_table(percent_path), 0, 1, target="responsivity").curve
    assert back.y_unit == "A/W"
    np.testing.assert_allclose(
        back.y_si,
        np.array([0.621, 0.382])
        * np.array([500e-9, 800e-9])
        * constants.e
        / constants.h
        / constants.c,
    )


def test_db_attenuation(tmp_path):
    path = write_text(
        tmp_path / "db.csv", ["Wavelength (nm),Loss (dB)", "500,0", "600,3.0103", "700,10"]
    )
    curve = import_curve(read_table(path), 0, 1).curve
    np.testing.assert_allclose(curve.y, [1.0, 0.5, 0.1], rtol=1e-4)


def test_percent_values_mislabeled_as_fraction_warn(tmp_path):
    path = write_text(tmp_path / "t.csv", ["Wavelength (nm),Transmission", "500,87.5", "600,95"])
    result = import_curve(read_table(path), 0, 1)
    assert len(result.warnings) == 1
    assert "outside [0, 1]" in result.warnings[0] and "percent" in result.warnings[0]
    fixed = import_curve(read_table(path), 0, 1, y_kind="percent")
    assert fixed.warnings == ()
    assert fixed.curve.y == [0.875, 0.95]


def test_rows_with_missing_values_are_dropped_with_a_note(tmp_path):
    path = write_text(
        tmp_path / "gap.csv",
        ["Wavelength (nm),Transmission (%)", "400,10", "500,n/a", "600,50", "700,60"],
    )
    result = import_curve(read_table(path), 0, 1)
    assert result.curve.x == [400.0, 600.0, 700.0]
    assert any("dropped 1 row" in w for w in result.warnings)


def test_unsorted_rows_are_sorted_and_duplicates_rejected(tmp_path):
    path = write_text(tmp_path / "u.csv", ["Wavelength (nm),T", "600,0.5", "400,0.1", "500,0.9"])
    assert import_curve(read_table(path), 0, 1).curve.x == [400.0, 500.0, 600.0]
    dup = write_text(tmp_path / "d.csv", ["Wavelength (nm),T", "400,0.1", "400,0.2", "500,0.9"])
    with pytest.raises(TableImportError, match="duplicate x value 400"):
        import_curve(read_table(dup), 0, 1)


def test_raw_kind_keeps_values_and_unit(tmp_path):
    path = write_text(tmp_path / "p.csv", ["Current (mA),Power (mW)", "10,0.5", "20,1.5", "30,2.0"])
    curve = import_curve(read_table(path), 0, 1).curve
    assert (curve.x_unit, curve.y_unit, curve.y, curve.x_label) == (
        "mA",
        "mW",
        [0.5, 1.5, 2.0],
        "current",
    )
    assert curve.y_si[0] == pytest.approx(0.5e-3)
    override = import_curve(read_table(path), 0, 1, y_label="power", y_unit="W").curve
    assert override.y_unit == "W"


def test_provenance_defaults_to_raw_data_with_filename(tmp_path):
    path = percent_csv(tmp_path)
    curve = import_curve(read_table(path), 0, 1).curve
    assert curve.provenance == Provenance(source="raw-data", reference="percent.csv")
    custom = Provenance(source="measured", by="Boran")
    assert import_curve(read_table(path), 0, 1, provenance=custom).curve.provenance == custom


def test_sheet_selection(tmp_path):
    wb = Workbook()
    first = wb.active
    first.title = "Notes"
    first.append(["nothing numeric here"])
    second = wb.create_sheet("TE")
    second.append(["Wavelength (nm)", "Transmission (%)"])
    second.append([400, 10])
    second.append([500, 20])
    path = tmp_path / "two.xlsx"
    wb.save(path)
    assert list_sheets(path) == ["Notes", "TE"]
    with pytest.raises(TableImportError, match="no numeric data"):
        read_table(path)
    assert import_curve(read_table(path, sheet="TE"), 0, 1).curve.y == [0.1, 0.2]
    with pytest.raises(TableImportError, match="no sheet 'TM'"):
        read_table(path, sheet="TM")


@pytest.mark.parametrize(
    ("name", "unit", "kind"),
    [
        ("Transmission", "%", "percent"),
        ("%T", None, "percent"),
        ("OD", None, "od"),
        ("Optical Density", None, "od"),
        ("Loss", "dB", "db"),
        ("Responsivity", "A/W", "responsivity"),
        ("Transmission", None, "fraction"),
        ("Power", "mW", "raw"),
    ],
)
def test_guess_y_kind(name, unit, kind):
    assert guess_y_kind(name, unit) == kind


@pytest.mark.parametrize(
    ("filename", "content", "message"),
    [
        ("missing.csv", None, "file not found"),
        ("empty.csv", "", "empty"),
        ("words.csv", "a,b\nc,d\n", "no numeric data"),
        ("one.csv", "x,y\n1,2\n", "at least 2 points"),
        ("data.pdf", "x", "unsupported file type"),
    ],
)
def test_unreadable_inputs_give_clear_errors(tmp_path, filename, content, message):
    path = tmp_path / filename
    if content is not None:
        path.write_text(content)
    with pytest.raises(TableImportError, match=message):
        import_curve(read_table(path), 0, 1, x_unit="nm")


def test_unknown_x_unit_and_negative_frequency(tmp_path):
    path = write_text(tmp_path / "a.csv", ["Wavelength (florps),T", "1,0.1", "2,0.2"])
    with pytest.raises(TableImportError, match="unknown unit"):
        import_curve(read_table(path), 0, 1)
    thz = write_text(tmp_path / "b.csv", ["Frequency (THz),T", "0,0.1", "2,0.2"])
    with pytest.raises(TableImportError, match="must be positive"):
        import_curve(read_table(thz), 0, 1)
    bad = write_text(tmp_path / "c.csv", ["Current (mA),Responsivity (A/W)", "1,0.1", "2,0.2"])
    with pytest.raises(TableImportError, match="needs a wavelength x axis"):
        import_curve(read_table(bad), 0, 1)
